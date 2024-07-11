import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

# Function to add checkboxes
def add_checkbox(ws, cell, checked):
    if checked:
        img = Image('checkbox_checked.png')  # Path to checked checkbox image
    else:
        img = Image('checkbox_unchecked.png')  # Path to unchecked checkbox image
    img.width, img.height = 15, 15  # Adjust the size of the checkbox image as needed
    img.anchor = cell
    ws.add_image(img)

# Read CSV files
csv_files = ['mvco_mlsl_qc.csv', 
            #  '../../../data/casper_west/RV_Sally_Ride_Flux/processed_csv/RVSR_2017_flux_10_v3_decorr.csv', 
             '../../../data/RVSR_2017_flux_10_v3_decorr.csv', 
             '../../../data/Flip_Flux_Processed_data.csv', 
             '../../../data/MAPS_20min_CASPER17_West.csv']
dataframes = [pd.read_csv(file) for file in csv_files]

# List of all variables
all_variables = set()
for df in dataframes:
    all_variables.update(df.columns)

# Convert the set of all variables to a list
all_variables = list(all_variables)

# Custom column names
column_names = ['MVCO', 'RVSR', 'FLIP', 'MAPS']

# Create a DataFrame to store the presence of variables
summary = pd.DataFrame(index=all_variables, columns=column_names)
summary = summary.fillna(False)

# Check for the presence of variables
for i, df in enumerate(dataframes):
    for variable in df.columns:
        summary.loc[variable, column_names[i]] = True

# Write to Excel
wb = Workbook()
ws = wb.active
ws.title = 'Summary'
ws.append(['Variable'] + column_names)

for row in summary.itertuples():
    ws.append([row.Index] + list(row[1:]))

# Add checkboxes
for i, row in enumerate(summary.itertuples(), start=2):
    for j in range(len(column_names)):
        checked = ws.cell(row=i, column=j+2).value  # +2 to account for the variable name column
        cell = f'{get_column_letter(j+2)}{i}'
        add_checkbox(ws, cell, checked)


wb.save('variable_summary.xlsx')
